# utils.py
"""Document ingestion helpers — extraction, cleaning, chunking, webhooks."""
import asyncio
import csv
import io
import os
import re
from typing import List

import httpx
from fastapi import HTTPException
from langchain.docstore.document import Document
from langchain.text_splitter import RecursiveCharacterTextSplitter

from config import settings


# ---------------------------------------------------------------------------
# Text cleaning (accuracy: removes noise that hurts embeddings + LLM)
# ---------------------------------------------------------------------------
_WS_RE = re.compile(r"[ \t]+")
_BLANK_RE = re.compile(r"\n{3,}")


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _WS_RE.sub(" ", text)
    text = _BLANK_RE.sub("\n\n", text)
    # drop control chars except newline/tab
    text = "".join(ch for ch in text if ch == "\n" or ch == "\t" or ord(ch) >= 32)
    return text.strip()


def _split(docs: List[Document]) -> List[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=getattr(settings, "CHUNK_SIZE", 900),
        chunk_overlap=getattr(settings, "CHUNK_OVERLAP", 150),
        separators=["\n\n", "\n", ". ", "? ", "! ", "; ", ": ", " ", ""],
        length_function=len,
    )
    chunks = splitter.split_documents(docs)
    # enrich metadata for better citations + dedup
    for idx, ch in enumerate(chunks):
        ch.metadata.setdefault("chunk_index", idx)
        ch.metadata["chunk_chars"] = len(ch.page_content)
    return chunks


# ---------------------------------------------------------------------------
# Extractors per file type
# ---------------------------------------------------------------------------
def _extract_pdf(content: bytes, filename: str) -> List[Document]:
    text_by_page: List[str] = []
    # Prefer pypdf (maintained fork of PyPDF2), fall back to PyPDF2
    reader = None
    last_err: Exception | None = None
    for mod_name in ("pypdf", "PyPDF2"):
        try:
            mod = __import__(mod_name)
            reader = mod.PdfReader(io.BytesIO(content))
            break
        except Exception as e:
            last_err = e
    if reader is None:
        raise HTTPException(status_code=500, detail=f"PDF library missing: {last_err}")

    try:
        if getattr(reader, "is_encrypted", False):
            try:
                reader.decrypt("")
            except Exception:
                raise HTTPException(status_code=400, detail=f"PDF '{filename}' is encrypted and cannot be read.")
        for i, page in enumerate(reader.pages):
            try:
                t = page.extract_text() or ""
            except Exception:
                t = ""
            t = clean_text(t)
            if t:
                text_by_page.append((i + 1, t))
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse PDF '{filename}': {e}")

    docs = [Document(page_content=t, metadata={"source": filename, "page": p, "type": "pdf"})
            for p, t in text_by_page]

    # Scanned PDF fallback: almost no text -> try OCR via pdf2image if available
    if not docs or sum(len(d.page_content) for d in docs) < 200:
        ocr_docs = _try_pdf_ocr(content, filename)
        if ocr_docs:
            return ocr_docs
    return docs


def _try_pdf_ocr(content: bytes, filename: str) -> List[Document]:
    try:
        from pdf2image import convert_from_bytes
        from PIL import Image as _Image  # noqa: F401
        import pytesseract
    except ImportError:
        return []
    try:
        pages = convert_from_bytes(content, dpi=200)
    except Exception:
        return []  # poppler missing etc. — caller already has (thin) text or will error
    docs: List[Document] = []
    import pytesseract as _pt
    for i, img in enumerate(pages[:20]):  # cap cost
        try:
            t = clean_text(_pt.image_to_string(img))
        except Exception:
            t = ""
        if len(t) > 20:
            docs.append(Document(page_content=t, metadata={"source": filename, "page": i + 1,
                                                            "type": "pdf-ocr"}))
    return docs


def _extract_docx(content: bytes, filename: str) -> List[Document]:
    try:
        import docx
    except ImportError:
        raise HTTPException(status_code=400, detail="DOCX support needs 'python-docx' (pip install python-docx).")
    try:
        doc = docx.Document(io.BytesIO(content))
        paras = [clean_text(p.text) for p in doc.paragraphs]
        # include tables (often hold the real data)
        for table in doc.tables:
            for row in table.rows:
                cells = [clean_text(c.text) for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    paras.append(line)
        text = "\n".join(p for p in paras if p)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse DOCX '{filename}': {e}")
    if not text.strip():
        raise HTTPException(status_code=400, detail=f"Could not extract any text from '{filename}'.")
    return [Document(page_content=text, metadata={"source": filename, "type": "docx"})]


def _extract_csv(content: bytes, filename: str) -> List[Document]:
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except Exception:
            continue
    else:
        raise HTTPException(status_code=400, detail=f"Could not decode CSV '{filename}'.")
    try:
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("no header row")
        lines = []
        for i, row in enumerate(reader):
            cells = [f"{k}: {v}" for k, v in row.items() if v and str(v).strip()]
            if cells:
                lines.append(f"Row {i + 1} — " + " | ".join(cells))
            if len(lines) >= 2000:
                break
        body = f"CSV file {filename}. Columns: {', '.join(reader.fieldnames)}.\n" + "\n".join(lines)
    except Exception:
        body = clean_text(text)[:200000]
    if not body.strip():
        raise HTTPException(status_code=400, detail=f"Could not extract any text from '{filename}'.")
    return [Document(page_content=body, metadata={"source": filename, "type": "csv"})]


def _extract_txt(content: bytes, filename: str) -> List[Document]:
    text = None
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            text = content.decode(enc)
            break
        except Exception:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail=f"Could not decode '{filename}'.")
    text = clean_text(text)
    if not text:
        raise HTTPException(status_code=400, detail=f"Could not extract any text from '{filename}'.")
    return [Document(page_content=text, metadata={"source": filename, "type": "text"})]


def _extract_image(content: bytes, filename: str) -> List[Document]:
    try:
        from PIL import Image
        import pytesseract
    except ImportError:
        raise HTTPException(status_code=400, detail="Image OCR needs Pillow + pytesseract installed.")
    try:
        image = Image.open(io.BytesIO(content))
        text = clean_text(pytesseract.image_to_string(image))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"OCR failed for '{filename}': {e}")
    if not text:
        raise HTTPException(status_code=400, detail=f"OCR found no text in '{filename}'.")
    return [Document(page_content=text, metadata={"source": filename, "type": "image-ocr"})]


def process_and_chunk_text(file_content: bytes, filename: str) -> List[Document]:
    """Extract -> clean -> chunk. Raises HTTPException(400) on bad input."""
    if not file_content:
        raise HTTPException(status_code=400, detail=f"File '{filename}' is empty.")
    max_bytes = int(getattr(settings, "MAX_FILE_MB", 25)) * 1024 * 1024
    if len(file_content) > max_bytes:
        raise HTTPException(status_code=413,
                            detail=f"File too large ({len(file_content) // 1024 // 1024}MB). Limit is {settings.MAX_FILE_MB}MB.")
    ext = os.path.splitext(filename)[1].lower()
    allowed = getattr(settings, "ALLOWED_EXT_SET", {".pdf", ".txt", ".md", ".csv", ".docx", ".png", ".jpg", ".jpeg"})
    if ext not in allowed:
        raise HTTPException(status_code=400,
                            detail=f"Unsupported file type '{ext}'. Allowed: {sorted(allowed)}")

    if ext == ".pdf":
        docs = _extract_pdf(file_content, filename)
    elif ext == ".docx":
        docs = _extract_docx(file_content, filename)
    elif ext == ".csv":
        docs = _extract_csv(file_content, filename)
    elif ext in (".txt", ".md"):
        docs = _extract_txt(file_content, filename)
    elif ext in (".png", ".jpg", ".jpeg"):
        docs = _extract_image(file_content, filename)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")

    if not docs:
        raise HTTPException(status_code=400, detail=f"Could not extract any text from '{filename}'.")
    chunks = _split(docs)
    # drop pathologically tiny chunks (headers/footers fragments)
    chunks = [c for c in chunks if len(c.page_content.strip()) >= 40] or chunks
    return chunks


def truncate_for_analysis(chunks: List[Document], max_chars: int = None) -> str:
    max_chars = max_chars or getattr(settings, "MAX_CONTEXT_CHARS", 12000)
    parts, total = [], 0
    for ch in chunks:
        t = ch.page_content
        if total + len(t) > max_chars:
            parts.append(t[: max_chars - total])
            break
        parts.append(t)
        total += len(t)
    return "\n\n".join(parts)


def format_docs(docs) -> str:
    blocks = []
    for d in docs:
        src = d.metadata.get("source", "unknown")
        page = d.metadata.get("page")
        tag = f"[Source: {src}" + (f", Page {page}" if page else "") + "]"
        blocks.append(f"{tag}\n{d.page_content}")
    return "\n\n".join(blocks)


async def trigger_n8n_webhooks(urls: List[str], data: dict):
    if not urls:
        return
    async with httpx.AsyncClient(timeout=10) as client:
        tasks = [client.post(url, json=data, timeout=10) for url in urls]
        results = await asyncio.gather(*tasks, return_exceptions=True)
        for i, res in enumerate(results):
            if isinstance(res, Exception):
                print(f"[n8n] ERROR {urls[i]}: {res}")
            else:
                print(f"[n8n] OK {urls[i]} -> {getattr(res, 'status_code', '?')}")
